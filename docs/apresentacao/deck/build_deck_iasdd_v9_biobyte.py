# -*- coding: utf-8 -*-
"""V9: sequência compacta e alinhada ao roteiro, com cinco vídeos curtos e um final."""
from pathlib import Path
import re
from pptx import Presentation

HERE = Path(__file__).resolve().parent
OUT = HERE / "output" / "apresentacao_iasdd_v9_biobyte.pptx"
import build_deck_iasdd_v7_biobyte as v7
from build_deck_iasdd_v7_biobyte import D, old, f

TOTAL = 99

def video_slide(prs, block, title, description, minutes, script):
    s = D._blank(prs)
    D.header(s, block, "Demonstração em vídeo", title)
    D._txt(s, .8, 2.0, 11.7, .8, [[(description, {"size": 22, "bold": True, "color": D.FAM["violet"][0]})]])
    D.cards(s, block, [
        ("O que será mostrado", description),
        ("Formato", "Jupyter Notebook ou ambiente LangNet, com as etapas visíveis e explicadas"),
        ("Duração", f"aproximadamente {minutes:g} minutos"),
    ], x=.8, y=3.25, size=16)
    D.footer(s, block, 0, TOTAL)
    old.v4.register(title, minutes, script)

def algorithm_slide(prs, title, purpose, strengths, health_use, minutes=0.7):
    s = D._blank(prs)
    D.header(s, 1, "Fundamentos de IA", title)
    D._txt(s, .75, 1.95, 11.5, .7, [[(purpose, {"size": 22, "bold": True, "color": D.FAM["violet"][0]})]])
    D.cards(s, 1, [("Como funciona", strengths), ("Onde é útil", health_use)], x=.75, y=3.15, size=18)
    D.footer(s, 1, 0, TOTAL)
    old.v4.register(title, minutes, f"Agora apresentamos {title}. {purpose} {strengths} Em saúde, {health_use}.")

def algorithms(prs):
    algorithm_slide(prs, "Regressão Logística", "Classifica casos e estima a probabilidade de cada classe.",
                    "Aprende uma fronteira entre classes e permite interpretar o efeito dos fatores.",
                    "triagem, risco e classificação binária quando explicabilidade é importante.")
    algorithm_slide(prs, "Árvore de Decisão", "Transforma decisões em regras condicionais legíveis.",
                    "Divide os dados por perguntas sucessivas e produz um caminho que pode ser auditado.",
                    "protocolos de triagem e regras clínicas que precisam ser explicadas ao usuário.")
    algorithm_slide(prs, "Random Forest", "Combina várias árvores para reduzir a instabilidade de uma árvore isolada.",
                    "Cada árvore aprende uma variação dos dados e a decisão final agrega seus resultados.",
                    "classificação robusta de risco quando há muitos fatores e relações não lineares.")
    algorithm_slide(prs, "Máquina de Vetores de Suporte", "Procura uma fronteira que separe as classes com maior margem.",
                    "Pode usar funções de núcleo para representar separações não lineares.",
                    "bases de tamanho moderado com muitas variáveis e separação complexa.")
    algorithm_slide(prs, "K Vizinhos Mais Próximos", "Classifica um caso pela semelhança com exemplos já conhecidos.",
                    "Consulta os vizinhos mais próximos; a escolha da escala e de K influencia o resultado.",
                    "protótipos didáticos, comparação por similaridade e o notebook usado na demonstração.")
    algorithm_slide(prs, "Redes Neurais e K-Means", "Redes neurais aprendem representações; K-Means agrupa sem rótulos.",
                    "A rede serve a padrões complexos; o K-Means procura grupos pela distância entre exemplos.",
                    "imagens e sinais para redes neurais; descoberta exploratória de perfis para agrupamento.")

def attention_history(prs):
    s = D._blank(prs)
    D.header(s, 2, "Modelos de linguagem", "De onde veio o Transformer?")
    D._txt(s, .75, 1.9, 11.6, .7, [[("Attention Is All You Need", {"size": 27, "bold": True, "color": D.FAM["violet"][0]})]])
    D._txt(s, .78, 2.65, 11.4, .45, [[("Vaswani et al. · Google Brain e University of Toronto · 2017", {"size": 16, "italic": True, "color": D.MUTED})]])
    D.cards(s, 2, [
        ("Antes", "RNNs e LSTMs processavam sequências passo a passo; dependências longas eram difíceis e o treinamento era pouco paralelizável."),
        ("A contribuição", "A atenção permite que cada token relacione-se diretamente com os demais, sem depender de uma recorrência."),
        ("O impacto", "Self-attention e treinamento paralelo tornaram viável escalar modelos de linguagem; a arquitetura passou a sustentar GPTs e sistemas atuais."),
    ], x=.75, y=3.25, size=15.5)
    D.footer(s, 2, 0, TOTAL)
    old.v4.register("De onde veio o Transformer?", .9,
        "Antes da anatomia, vale registrar a origem. O artigo Attention Is All You Need, de Ashish Vaswani e colaboradores, publicado em 2017 por pesquisadores do Google Brain e da University of Toronto, propôs substituir a recorrência por atenção. A ideia central é permitir que cada token considere diretamente os demais tokens da sequência. Isso tornou o treinamento muito mais paralelizável e abriu o caminho para escalar modelos de linguagem. A referência completa ficará nas fontes da apresentação.")

def environment_slide(prs, title, subtitle, code, explanation, script):
    s = D._blank(prs)
    D.header(s, 8, "Ambientes de geração de código", title)
    D._txt(s, .7, 1.82, 11.8, .6, [[(subtitle, {"size": 18, "bold": True, "color": D.FAM["violet"][0]})]])
    D.code_block(s, code, x=.75, y=2.65, w=6.4, size=11)
    D.callout(s, 8, [(explanation, {"size": 15, "color": D.INK})], x=7.45, y=2.65, w=5.0, h=3.25)
    D.footer(s, 8, 0, TOTAL)
    old.v4.register(title, .8, script)

def environment_overview(prs):
    s = D._blank(prs)
    D.header(s, 8, "Ambientes de geração de código", "Onde o agente trabalha?")
    D.cards(s, 8, [
        ("Claude Code", "agente de terminal orientado ao repositório; lê, altera, testa e pode conectar MCP"),
        ("Codex", "agente da OpenAI para explorar, implementar, revisar e automatizar tarefas com permissões controladas"),
        ("Cursor", "editor com agente integrado, contexto do projeto, edição inline e execução assistida"),
    ], x=.75, y=2.0, size=16)
    D.callout(s, 8, [("A diferença não é só a interface: é como cada ambiente organiza contexto, ferramentas, permissões, sessões e evidência.", {"size": 17, "bold": True, "color": D.INK})], y=5.8, h=.75)
    D.footer(s, 8, 0, TOTAL)
    old.v4.register("Onde o agente trabalha?", .8,
        "Antes dos protocolos, precisamos localizar os ambientes de desenvolvimento. Claude Code, Codex e Cursor não são apenas caixas de chat: eles operam sobre um repositório, usam contexto, chamam ferramentas, editam arquivos e precisam deixar mudanças verificáveis. A escolha do ambiente deve considerar permissões, rastreabilidade, integração com o editor, automação e o modo como a equipe revisa o resultado.")

def environment_commands(prs):
    environment_slide(prs, "Claude Code: comandos e extensões", "O terminal vira uma sessão de engenharia com contexto persistente.",
        "claude\nclaude -c\nclaude -r <session-id>\nclaude -p \"run the tests\"\nclaude mcp\n/help   /compact   /resume",
        "Skills codificam fluxos repetíveis; subagentes e equipes de agentes dividem investigação e implementação; MCP conecta ferramentas externas. O ambiente também permite continuar ou retomar sessões e executar em modo não interativo.",
        "Este é o Claude Code, o ambiente que estamos chamando de Codecodes. O comando inicia uma sessão no repositório. É possível continuar a conversa mais recente, retomar uma sessão pelo identificador, executar uma solicitação não interativa e configurar servidores MCP. Skills transformam procedimentos em instruções reutilizáveis. Em fluxos mais complexos, subagentes ou equipes podem dividir investigação, implementação e revisão; esses recursos devem sempre manter limites e evidências claras.")

def codex_commands(prs):
    environment_slide(prs, "Codex: execução, revisão e skills", "Um agente de código pode ser usado interativamente ou em automações reproduzíveis.",
        "codex\ncodex exec \"inspect the repository\"\ncodex exec --json \"run the checks\"\ncodex exec --output-schema schema.json \"evaluate\"\n/init\n$skill-creator",
        "A sessão interativa serve à exploração. codex exec serve a automação; --json registra eventos e --output-schema produz uma saída verificável. Skills e AGENTS.md transformam convenções em contexto reutilizável.",
        "No Codex, podemos trabalhar interativamente ou usar codex exec em tarefas automatizadas. O modo JSON permite registrar eventos e analisar o que aconteceu; um schema de saída transforma a resposta em um artefato verificável. O comando init ajuda a criar as instruções do repositório, e uma skill organiza um procedimento repetível. O ponto central é o mesmo do SDD: explicitar contexto, permissões, critérios e validação antes de delegar.")

def cursor_commands(prs):
    environment_slide(prs, "Cursor e editores com agentes", "A mesma capacidade agêntica aparece dentro do ciclo de edição.",
        "selecionar contexto\n→ pedir alteração\n→ revisar diff\n→ executar testes\n→ aceitar ou rejeitar",
        "O editor aproxima contexto, código, diff, terminal e revisão. É produtivo para iterações rápidas, mas a aprovação humana e os gates do projeto continuam necessários.",
        "O Cursor representa a família de editores com agentes integrados. A vantagem é aproximar seleção de contexto, edição, diff, terminal e revisão. Isso reduz a troca de ferramentas, mas não substitui o processo: cada alteração ainda precisa de testes, revisão e critérios de aceitação. O ambiente acelera a execução; a especificação continua definindo o que deve ser executado.")

def divider(prs, block, number, title, subtitle, message):
    old.divider(prs, block, number, title, subtitle, message)

def strip_internal_block_labels(path):
    prs = Presentation(str(path))
    for slide in prs.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.text = re.sub(r"^BLOCO\s+\d+\s*[·•]\s*", "", run.text, flags=re.I)
    prs.save(str(path))

def build():
    f.META = []
    prs = f.D.new_prs()
    funcs = [
        f.s1, f.s2, f.s3,
        lambda p: divider(p, 1, 1, "Introdução e Machine Learning", "histórico, pipeline e algoritmos", "Começamos pelo contexto mínimo e pelo fluxo completo de aprendizagem em dados."),
        f.s4, old.v4.ml_pipeline, old.v4.ml_evaluation, algorithms,
        lambda p: video_slide(p, 1, "Demonstração: pipeline de Machine Learning", "Do carregamento da base à avaliação do classificador.", 8,
            "Agora veremos, em oito minutos, o notebook executando o pipeline completo: preparação, visualização, divisão, treinamento e avaliação."),
        lambda p: divider(p, 2, 2, "IA na saúde", "prognóstico, diagnóstico e tratamento", "Aplicamos os conceitos a problemas clínicos sem misturar tarefas diferentes."),
        f.s5, f.s5b, f.s5c,
        lambda p: video_slide(p, 2, "Demonstração: IA na saúde", "Um exemplo de classificação e apoio à decisão em saúde.", 8,
            "Neste vídeo, veremos um exemplo curto de aplicação dos algoritmos em saúde, observando entradas, saída, avaliação e limitações."),
        lambda p: divider(p, 3, 3, "Deep Learning e Transfer Learning", "redes neurais e adaptação de modelos", "Depois do aprendizado clássico, mostramos como redes profundas reutilizam conhecimento."),
        v7.transfer,
        lambda p: video_slide(p, 3, "Demonstração: treinamento e Transfer Learning", "Da rede pré-treinada ao ajuste de uma tarefa específica.", 8,
            "O vídeo mostra a preparação de uma rede, o congelamento das camadas gerais e o ajuste da camada específica."),
        lambda p: divider(p, 4, 4, "Modelos de linguagem e Transformer", "treinamento, atenção e geração", "Agora passamos das imagens e tabelas para modelos que aprendem sequências de tokens."),
        old.llm_intro, attention_history, f.s6, f.s7, f.s8, f.s9, f.s10, f.s11, f.s12, old.finetune_intro,
        lambda p: video_slide(p, 4, "Demonstração: estrutura do Transformer", "Uma visualização do fluxo de tokens, atenção e geração.", 8,
            "Neste vídeo, acompanhamos a estrutura do Transformer e relacionamos treinamento, atenção e geração de tokens."),
        lambda p: divider(p, 5, 5, "Agentes, contexto e execução", "entrada, tarefa, memória, tools e composição", "Um agente é definido pela tarefa que recebe e pelo resultado que devolve."),
        v7.agent_definition, old.agent_intro, f.s21, f.s22b, f.s23, f.s24, f.s25, f.s26, f.s27, f.s28,
        lambda p: video_slide(p, 5, "Demonstração: sistema multiagente", "Agentes especializados coordenados em uma tarefa do BioByte.", 8,
            "Agora veremos agentes especializados trocando contexto, usando ferramentas e devolvendo resultados verificáveis."),
        lambda p: divider(p, 6, 6, "Contexto, RAG e tecnologias de agentes", "LangChain, LangGraph, CrewAI, AutoGen e Qdrant", "Depois de definir o agente, mostramos como construir e alimentar sua execução."),
        old.context_intro, f.s16, v7.rag_code, f.s17, f.s18, f.s19, f.s20,
        lambda p: divider(p, 7, 7, "Frameworks de agentes", "composição e controle da execução", "Primeiro comparamos as bibliotecas; depois veremos os ambientes onde elas são utilizadas."),
        f.s34, f.s35, f.s36, f.s37, f.s38, f.s39, f.s40,
        lambda p: divider(p, 8, 8, "Ambientes de geração de código", "Claude Code, Codex e Cursor", "Agora saímos do framework e observamos o ambiente de trabalho do desenvolvedor."),
        environment_overview, environment_commands, codex_commands, cursor_commands, f.s41, f.s42, f.s43,
        lambda p: divider(p, 9, 9, "Protocolos e contratos", "MCP, A2A e formatos de interoperabilidade", "Com os ambientes definidos, podemos discutir como os componentes se conectam."),
        f.s29, f.s30, f.s31, f.s32, f.s33,
        lambda p: divider(p, 10, 10, "SDD: especificar antes de delegar", "da intenção ao artefato verificável", "A especificação conecta requisitos, agentes, código, testes e rastreabilidade."),
        f.s44, f.s45, f.s46, old.sdd_intro, f.s47, f.s48, f.s49, f.s50, f.s51, f.s52, f.s53, f.s54, f.s55,
        lambda p: divider(p, 11, 11, "BioByte e LangNet", "a cadeia completa de especificação à execução", "Agora conectamos o método ao sistema real."),
        old.biobyte_intro, old.v4.base.s60, old.v4.base.s61, f.s62, f.s63, f.s64, f.s65, f.s66,
        lambda p: video_slide(p, 11, "Demonstração final: BioByte no LangNet", "Da especificação aos agentes, tarefas, rede de Petri, código, testes e aplicação.", 20,
            "Agora começa a demonstração final de vinte minutos. Ela percorre o pipeline do LangNet até a aplicação BioByte executável e verificável."),
        lambda p: divider(p, 12, 12, "Fechamento", "método, evidência e próximos passos", "Encerramos retomando a tese e abrindo para perguntas."), f.s67,
    ]
    for fn in funcs:
        fn(prs)
    # Duração global: 58 minutos de lâminas + cinco vídeos de 8 minutos
    # + demonstração final de 20 minutos.
    video_rows = [x for x in f.META if x["title"].startswith("Demonstração")]
    for x in video_rows:
        x["minutes"] = 20.0 if x["title"].startswith("Demonstração final") else 8.0
    spoken_rows = [x for x in f.META if x not in video_rows]
    factor = 58.0 / sum(x["minutes"] for x in spoken_rows)
    for x in spoken_rows:
        x["minutes"] *= factor
    acum = 0
    for x in f.META:
        acum += x["minutes"]; x["acum"] = acum
    prs.save(str(OUT))
    strip_internal_block_labels(OUT)
    old.v4.fix_text_and_numbers(OUT)
    print(f"PPTX: {OUT} slides: {len(prs.slides)} slides_minutes: {sum(x['minutes'] for x in f.META):.2f}")

if __name__ == "__main__": build()

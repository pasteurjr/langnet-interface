# -*- coding: utf-8 -*-
"""V15 — NÃO sobrescreve a V14. A V14 + v15_mudancas + v15_falas."""
from pathlib import Path
import sys

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import decklib as D
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

import build_deck_iasdd_v10_biobyte as V10
import build_deck_iasdd_v11_biobyte as V11
import build_deck_iasdd_v12_biobyte as V12
import build_deck_iasdd_v13_biobyte as V13
import build_deck_iasdd_v14_biobyte as V14
import v15_mudancas
import v15_falas

OUT = HERE / "output" / "apresentacao_iasdd_v15_biobyte.pptx"

SLIDES = v15_falas.aplicar(v15_mudancas.aplicar(V14.SLIDES))
TOTAL = len(SLIDES)
V10.TOTAL = V11.TOTAL = V12.TOTAL = V13.TOTAL = V14.TOTAL = TOTAL
V14.SLIDES = SLIDES   # o sumário lê os blocos daqui


def _algoritmo_fig(prs, s, n):
    """ideia central + três cartões à esquerda + a figura da apresentação-base à direita"""
    sl = D._blank(prs); b = s["bloco"]
    c1, c2, c3 = D.header(sl, b, s["chapeu"], s["titulo"])
    D._rect(sl, 0.6, 1.98, D.SW - 1.2, 0.7, fill=c3, rounded=True, radius=0.08)
    D._rect(sl, 0.6, 1.98, 0.12, 0.7, fill=D.FAM["amber"][0])
    D._txt(sl, 0.95, 1.98, D.SW - 1.6, 0.7,
           [[("Ideia central   ", {"size": 14, "bold": True, "color": D.MUTED}),
             (s["ideia"], {"size": 18, "bold": True, "color": c1})]], anchor=MSO_ANCHOR.MIDDLE)
    from PIL import Image
    iw, ih = Image.open(HERE / s["imagem"]).size
    if iw / ih > 1.9:
        # figura larga: ocupa a largura toda, e os três cartões ficam lado a lado embaixo
        hc = 1.55
        yc = V11.BAIXO - hc
        D.image_center(sl, str(HERE / s["imagem"]), x=0.6, y=2.8, w=D.SW - 1.2, h=yc - 2.8 - 0.12)
        gap = 0.18
        w = (D.SW - 1.2 - 2 * gap) / 3
        for k, (lead, txt) in enumerate([("Como funciona", s["como"]), ("Para que se presta", s["serve"]),
                                         ("Onde é mais adequado", s["onde"])]):
            x = 0.6 + k * (w + gap)
            D._rect(sl, x, yc, w, hc, fill=c3, rounded=True, radius=0.08)
            D._rect(sl, x, yc, 0.1, hc, fill=c1)
            D._txt(sl, x + 0.25, yc + 0.05, w - 0.4, hc - 0.1,
                   [[(lead, {"size": 13, "bold": True, "color": c1})], [(txt, {"size": 11.5, "color": D.INK})]],
                   anchor=MSO_ANCHOR.MIDDLE, space_after=2, line_spacing=1.02)
        D.footer(sl, b, n, TOTAL); D.notes(sl, s.get("fala", ""))
        return sl
    wl = 5.9
    D.cards(sl, b, [("Como funciona", s["como"]), ("Para que se presta", s["serve"]),
                    ("Onde é mais adequado", s["onde"])], x=0.6, y=2.85, w=wl, size=12.5, ch=1.33)
    D.image_center(sl, str(HERE / s["imagem"]), x=0.6 + wl + 0.25, y=2.85, w=D.SW - 1.2 - wl - 0.25,
                   h=V11.BAIXO - 2.85)
    D.footer(sl, b, n, TOTAL); D.notes(sl, s.get("fala", ""))
    return sl


def _figuras2(prs, s, n):
    """duas figuras empilhadas, cada uma com um rótulo à esquerda"""
    sl = D._blank(prs); b = s["bloco"]
    c1, c2, c3 = D.header(sl, b, s["chapeu"], s["titulo"])
    y = 2.0
    for rot, img, h in s["imagens"]:
        D._rect(sl, 0.6, y, 0.08, h, fill=c1)
        D._txt(sl, 0.8, y, 1.6, h, [[(rot, {"size": 14, "bold": True, "color": c1})]], anchor=MSO_ANCHOR.MIDDLE)
        D.image_center(sl, str(HERE / img), x=2.45, y=y, w=D.SW - 0.6 - 2.45, h=h)
        y += h + 0.2
    D.footer(sl, b, n, TOTAL); D.notes(sl, s.get("fala", ""))
    return sl


def _tab(sl, b, colunas, linhas, larguras, y, fs, rh, col_forte=None):
    from pptx.util import Inches, Pt
    c1, c2, c3 = D.fam(b)
    nr = len(linhas) + 1
    gf = sl.shapes.add_table(nr, len(colunas), Inches(0.6), Inches(y), Inches(sum(larguras)), Inches(rh * nr))
    t = gf.table
    for k, w in enumerate(larguras):
        t.columns[k].width = Inches(w)
    for r in range(nr):
        t.rows[r].height = Inches(rh)
    for r, linha in enumerate([colunas] + linhas):
        for k, val in enumerate(linha):
            c = t.cell(r, k); c.fill.solid()
            c.fill.fore_color.rgb = c1 if r == 0 else (c3 if r % 2 else D.WHITE)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.margin_left = c.margin_right = Inches(0.07); c.margin_top = c.margin_bottom = Inches(0.01)
            p = c.text_frame.paragraphs[0]; c.text_frame.word_wrap = True
            forte = r == 0 or val.startswith("*") or (k == col_forte and r > 0)
            run = p.add_run(); run.text = val.lstrip("*")
            run.font.name = D.FONT; run.font.size = Pt(fs); run.font.bold = forte
            run.font.color.rgb = D.WHITE if r == 0 else (c1 if forte else D.INK)
    return y + rh * nr


def _tabela2(prs, s, n):
    """tabela dos modelos (compacta) + tabela de comparação de desempenho por tarefa"""
    sl = D._blank(prs); b = s["bloco"]
    c1, c2, c3 = D.header(sl, b, s["chapeu"], s["titulo"])
    larg = s.get("larguras") or [12.13 / len(s["colunas"])] * len(s["colunas"])
    y = _tab(sl, b, s["colunas"], s["linhas"], larg, 1.98, 10.5, 0.27)
    cp = s["comparacao"]
    y += 0.1
    D._txt(sl, 0.6, y, 8, 0.3, [[(cp["titulo"], {"size": 14, "bold": True, "color": c1})]])
    y = _tab(sl, b, cp["colunas"], cp["linhas"], cp["larguras"], y + 0.33, 10.5, 0.27, col_forte=1)
    D._txt(sl, 0.6, V11.BAIXO - 0.25, D.SW - 1.2, 0.25, [[(cp["nota"], {"size": 9.5, "italic": True, "color": D.MUTED})]])
    D.footer(sl, b, n, TOTAL); D.notes(sl, s.get("fala", ""))
    return sl


DESENHO = dict(V14.DESENHO)
DESENHO["tabela2"] = _tabela2
DESENHO["algoritmo_fig"] = _algoritmo_fig
DESENHO["figuras2"] = _figuras2


def build():
    prs = D.new_prs()
    for i, s in enumerate(SLIDES, start=1):
        fn = DESENHO[s["tipo"]]
        if s["tipo"] == "capa":
            fn(prs, s)
        else:
            fn(prs, s, i)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    caminho = build()
    palavras = sum(len(s.get("fala", "").split()) for s in SLIDES)
    video = sum(s.get("minutos", 0) for s in SLIDES if s["tipo"] == "video")
    print("gravado: %s" % caminho)
    print("lâminas: %d" % TOTAL)
    print("palavras de narração: %d  (%.1f min a 135 palavras/min)" % (palavras, palavras / 135.0))
    print("vídeos: %g min" % video)

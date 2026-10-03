# -*- coding: utf-8 -*-
"""Gera a V3 do deck, com a fonte da seção de IA na saúde identificada."""
from pathlib import Path
import shutil

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

HERE = Path(__file__).resolve().parent
V2 = HERE / "output" / "apresentacao_iasdd_v2_biobyte.pptx"
OUT = HERE / "output" / "apresentacao_iasdd_v3_biobyte.pptx"


def build():
    # A V3 nasce da V2; a V2 original não é sobrescrita.
    if not V2.exists():
        import build_deck_iasdd_v2_biobyte as source
        source.build_v2()
    shutil.copy2(V2, OUT)
    prs = Presentation(str(OUT))
    source_by_slide = {
        5: "Fonte: Algoritmos Inteligentes3.pptx, slides 155–237 (prognóstico e risco).",
        6: "Fonte: Algoritmos Inteligentes3.pptx, slides 238–284 (diagnóstico e imagens médicas).",
        7: "Fonte: Algoritmos Inteligentes3.pptx, slides 285–339 (tratamento e inferência causal).",
    }
    for number, source_text in source_by_slide.items():
        slide = prs.slides[number - 1]
        # Corrige a referência imprecisa à origem do exemplo de COVID-19.
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.text = run.text.replace(
                        "o caso do PDF: COVID-19 em tomografia",
                        "o exemplo de COVID-19 em tomografia da fonte original",
                    )
        box = slide.shapes.add_textbox(Inches(0.65), Inches(6.78), Inches(10.8), Inches(0.22))
        tf = box.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = source_text
        p.font.name = "Aptos"
        p.font.size = Pt(7.5)
        p.font.italic = True
        p.font.color.rgb = RGBColor(92, 102, 126)
    prs.save(str(OUT))
    print(f"PPTX: {OUT} slides: {len(prs.slides)}")


if __name__ == "__main__":
    build()

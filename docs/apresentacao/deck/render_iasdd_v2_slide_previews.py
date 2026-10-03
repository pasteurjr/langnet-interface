# -*- coding: utf-8 -*-
"""Renderizador de previews visuais do PPTX para o roteiro narrado."""
from pathlib import Path
import textwrap

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parent
PPTX = ROOT / "output" / "apresentacao_iasdd_v2_biobyte.pptx"
OUT = ROOT.parent / "roteiro_narrado_iasdd_v2_biobyte_slides"
W, H = 1600, 900


def font(size, bold=False, mono=False):
    names = []
    if mono:
        names = ["/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"]
    elif bold:
        names = ["/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf"]
    else:
        names = ["/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"]
    for name in names:
        if Path(name).exists():
            return ImageFont.truetype(name, max(10, int(size * 1.35)))
    return ImageFont.load_default()


def rgb(value, default=(22, 30, 51)):
    try:
        return tuple(value.rgb)
    except Exception:
        return default


def text_color(shape):
    try:
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                return rgb(r.font.color)
    except Exception:
        pass
    return (22, 30, 51)


def draw_text(draw, box, text, size=20, color=(22, 30, 51), bold=False, mono=False, align="left"):
    x, y, w, h = box
    f = font(size, bold=bold, mono=mono)
    words = text.replace("\n", " \n ").split()
    lines, current = [], ""
    max_chars = max(8, int(w / max(7, size * .62)))
    for word in words:
        if word == "\n":
            lines.append(current); current = ""; continue
        trial = (current + " " + word).strip()
        if len(trial) > max_chars and current:
            lines.append(current); current = word
        else:
            current = trial
    if current: lines.append(current)
    line_h = max(14, int(size * 1.35))
    max_lines = max(1, int(h / line_h))
    lines = lines[:max_lines]
    for i, line in enumerate(lines):
        yy = y + i * line_h
        if align == "center":
            bb = draw.textbbox((0, 0), line, font=f)
            xx = x + max(0, (w - (bb[2] - bb[0])) // 2)
        elif align == "right":
            bb = draw.textbbox((0, 0), line, font=f)
            xx = x + max(0, w - (bb[2] - bb[0]))
        else:
            xx = x
        draw.text((xx, yy), line, fill=color, font=f)


def shape_box(shape, sx, sy):
    return (int(shape.left * sx), int(shape.top * sy), int(shape.width * sx), int(shape.height * sy))


def fill_color(shape, default=(255, 255, 255)):
    try:
        if shape.fill.type is not None:
            return rgb(shape.fill.fore_color, default)
    except Exception:
        pass
    return default


def render_slide(slide, number, total):
    sx, sy = W / 12191695, H / 6858000
    img = Image.new("RGB", (W, H), (251, 252, 254))
    draw = ImageDraw.Draw(img)

    for shape in slide.shapes:
        x, y, w, h = shape_box(shape, sx, sy)
        if w <= 0 or h <= 0:
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            try:
                from io import BytesIO
                pic = Image.open(BytesIO(shape.image.blob)).convert("RGB")
                pic.thumbnail((w, h))
                px = x + (w - pic.width) // 2
                py = y + (h - pic.height) // 2
                img.paste(pic, (px, py))
            except Exception:
                pass
            continue
        if getattr(shape, "has_table", False):
            table = shape.table
            rh = max(1, h // len(table.rows))
            cw = max(1, w // len(table.columns))
            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells):
                    cx, cy = x + ci * cw, y + ri * rh
                    draw.rectangle((cx, cy, cx + cw, cy + rh), fill=(228, 237, 248) if ri == 0 else (255, 255, 255), outline=(170, 180, 195))
                    draw_text(draw, (cx + 7, cy + 4, cw - 14, rh - 8), cell.text, size=10 if ri else 11, color=(255,255,255) if ri == 0 else (22,30,51), bold=ri == 0)
            continue
        if getattr(shape, "has_text_frame", False):
            try:
                fill = fill_color(shape, (251, 252, 254))
                if fill != (251, 252, 254) or shape.shape_type != MSO_SHAPE_TYPE.TEXT_BOX:
                    draw.rectangle((x, y, x+w, y+h), fill=fill, outline=(225, 231, 240))
            except Exception:
                pass
            text = shape.text.strip()
            if not text:
                continue
            size = 18
            bold = False
            mono = False
            try:
                p = shape.text_frame.paragraphs[0]
                if p.runs:
                    run = p.runs[0]
                    size = int((run.font.size.pt if run.font.size else 18))
                    bold = bool(run.font.bold)
                    mono = run.font.name and "Mono" in run.font.name
            except Exception:
                pass
            if " / " in text:
                size = 12
            draw_text(draw, (x + 10, y + 7, max(20, w - 20), max(20, h - 14)), text, size=size, color=text_color(shape), bold=bold, mono=mono)
    draw.text((W - 130, 18), f"{number:02d}/{total:02d}", fill=(90, 102, 126), font=font(14, bold=True))
    return img


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    prs = Presentation(str(PPTX))
    total = len(prs.slides)
    for i, slide in enumerate(prs.slides, start=1):
        render_slide(slide, i, total).save(OUT / f"slide_{i:02d}.png", optimize=True)
    print("renders:", len(prs.slides), "dir:", OUT)


if __name__ == "__main__":
    main()
